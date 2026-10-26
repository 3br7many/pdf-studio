#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
PDF Studio — a fully local PDF toolkit with first-class Arabic support.
Nothing ever leaves this machine; the server binds to 127.0.0.1 only.
"""

import io
import os
import re
import sys
import hashlib
import shutil
import zipfile
import tempfile
import mimetypes
import threading
import unicodedata
import webbrowser
from datetime import datetime

from flask import Flask, request, jsonify, send_file, render_template, abort

import pypdf
from pypdf import PdfReader, PdfWriter
import pikepdf
import pymupdf  # PyMuPDF
import pdfplumber
from PIL import Image
from reportlab.pdfgen import canvas as rl_canvas
from reportlab.lib.colors import HexColor
from reportlab.lib.utils import ImageReader
from docx import Document as DocxDocument
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from openpyxl import Workbook
from openpyxl.styles import Alignment

def _search_dirs():
    """Where to look for templates/ and fonts/.

    Three cases have to work: running from source, running from the packaged
    application (PyInstaller unpacks its data to a temporary folder), and
    running an applied update (whose code sits in the user's data folder and
    carries the templates but not the fonts). Searching in order covers all
    three without special cases."""
    dirs = [os.path.dirname(os.path.abspath(__file__))]
    home = os.environ.get("PDFSTUDIO_HOME")     # set by the launcher
    if home:
        dirs.append(home)
    if getattr(sys, "frozen", False):
        meipass = getattr(sys, "_MEIPASS", None)
        if meipass:
            dirs.append(meipass)
        dirs.append(os.path.dirname(sys.executable))
    return dirs


def resource(name):
    for d in _search_dirs():
        candidate = os.path.join(d, name)
        if os.path.exists(candidate):
            return candidate
    return os.path.join(_search_dirs()[-1], name)


APP_DIR = os.path.dirname(os.path.abspath(__file__))
FONT_DIR = resource("fonts")
FROZEN = bool(getattr(sys, "frozen", False))
WORK = os.path.join(tempfile.gettempdir(), "pdfstudio_work")
os.makedirs(WORK, exist_ok=True)

app = Flask(__name__, template_folder=resource("templates"))
app.config["MAX_CONTENT_LENGTH"] = 512 * 1024 * 1024  # 512 MB

FILES = {}
_lock = threading.Lock()
_counter = {"n": 0}


# ============================================================================
# Helpers
# ============================================================================
def new_fid(prefix="f"):
    with _lock:
        _counter["n"] += 1
        return f"{prefix}{_counter['n']}_{os.urandom(3).hex()}"


def safe_name(name):
    """Keeps Arabic/Unicode characters, strips only path separators and
    characters that are illegal in filenames on Windows."""
    name = os.path.basename(name or "file")
    name = re.sub(r'[\\/:*?"<>|\r\n\t]+', "_", name).strip() or "file"
    return name[:120]


def stem(name):
    return os.path.splitext(name)[0]


def register(path, name, kind="pdf", origin="upload"):
    fid = new_fid("o" if origin == "output" else "f")
    info = {
        "id": fid,
        "name": name,
        "path": path,
        "kind": kind,
        "origin": origin,
        "size": os.path.getsize(path),
        "created": datetime.now().strftime("%H:%M:%S"),
        "pages": 0,
        "encrypted": False,
        "unlocked": True,
        "has_form": False,
    }
    if kind == "pdf":
        info.update(probe_pdf(path))
    FILES[fid] = info
    return public(info)


def probe_pdf(path):
    out = {"pages": 0, "encrypted": False, "unlocked": True, "has_form": False}
    try:
        r = PdfReader(path)
        if r.is_encrypted:
            out["encrypted"] = True
            try:
                ok = r.decrypt("")
            except Exception:
                ok = 0
            if not ok:
                out["unlocked"] = False
                return out
        out["pages"] = len(r.pages)
        try:
            out["has_form"] = bool(r.get_fields())
        except Exception:
            out["has_form"] = False
    except Exception:
        pass
    return out


def public(f):
    d = {k: v for k, v in f.items() if k != "path"}
    d["size_h"] = human(f["size"])
    return d


def human(n):
    for u in ["B", "KB", "MB", "GB"]:
        if n < 1024:
            return f"{n:.0f} {u}" if u == "B" else f"{n:.1f} {u}"
        n /= 1024
    return f"{n:.1f} TB"


def get_file(fid, kind=None):
    info = FILES.get(fid)
    if not info or not os.path.exists(info["path"]):
        abort(404, description="File not found — please upload it again.")
    if kind and info["kind"] != kind:
        abort(400, description="That file type can't be used for this operation.")
    if info["kind"] == "pdf" and not info.get("unlocked", True):
        abort(400, description="This PDF is password-protected — unlock it first "
                               "from the Security tab.")
    return info


def out_path(name):
    return os.path.join(WORK, f"{new_fid('out')}__{safe_name(name)}")


def parse_pages(spec, total):
    """'1-3,5,8-' -> [0,1,2,4,7,...]  (1-based in, 0-based out)"""
    if not spec or str(spec).strip().lower() in ("all", "*"):
        return list(range(total))
    result = []
    for part in str(spec).replace("،", ",").split(","):
        part = part.strip()
        if not part:
            continue
        m = re.match(r"^(\d+)\s*-\s*(\d*)$", part)
        if m:
            a = int(m.group(1))
            b = int(m.group(2)) if m.group(2) else total
            for p in range(a, b + 1):
                if 1 <= p <= total:
                    result.append(p - 1)
        elif part.isdigit():
            p = int(part)
            if 1 <= p <= total:
                result.append(p - 1)
    seen, uniq = set(), []
    for p in result:
        if p not in seen:
            seen.add(p)
            uniq.append(p)
    return uniq or list(range(total))


def reader_of(info):
    r = PdfReader(info["path"])
    if r.is_encrypted:
        r.decrypt("")
    return r


def finish(writer, name, note=None):
    p = out_path(name)
    with open(p, "wb") as fh:
        writer.write(fh)
    payload = {"ok": True, "file": register(p, safe_name(name), "pdf", "output")}
    if note:
        payload["note"] = note
    return jsonify(**payload)


def err(msg, code=400):
    return jsonify(ok=False, error=str(msg)), code


@app.errorhandler(400)
@app.errorhandler(404)
@app.errorhandler(413)
@app.errorhandler(500)
def handle_err(e):
    desc = getattr(e, "description", str(e))
    code = getattr(e, "code", 500)
    if request.path.startswith("/api/"):
        return jsonify(ok=False, error=desc), code
    return desc, code


def hex_rgb(color, default=(0, 0, 0)):
    c = str(color or "").strip().lstrip("#")
    if len(c) == 3:
        c = "".join(ch * 2 for ch in c)
    if len(c) != 6:
        return default
    try:
        return tuple(int(c[i:i + 2], 16) / 255 for i in (0, 2, 4))
    except ValueError:
        return default


def hex_norm(color, default="#000000"):
    c = str(color or "").strip().lstrip("#")
    if len(c) == 3:
        c = "".join(ch * 2 for ch in c)
    return "#" + c if re.fullmatch(r"[0-9a-fA-F]{6}", c or "") else default


# ============================================================================
# Arabic support
# ============================================================================
AR_RANGE = r"\u0600-\u06FF\u0750-\u077F\u08A0-\u08FF\uFB50-\uFDFF\uFE70-\uFEFF"
RE_ARABIC = re.compile(f"[{AR_RANGE}]")
# Presentation forms: the "already-shaped" glyphs many PDFs store instead of
# real letters. Their presence means the text is in visual (rendered) order.
RE_PRESENTATION = re.compile(r"[\uFB50-\uFDFF\uFE70-\uFEFF]")
# A whole left-to-right run, spaces included, so that "Notes Final" flips back
# to "Final Notes" instead of each word being reversed in place. Punctuation is
# only allowed *between* alphanumerics, which keeps "120,000 150,000" together
# as one run instead of splitting it at the comma.
_LTR_TOKEN = r"[A-Za-z0-9]+(?:[@._\-/:%#&+=,'’()][A-Za-z0-9]+)*"
RE_LTR_RUN = re.compile(rf"{_LTR_TOKEN}(?:[ \t]+{_LTR_TOKEN})*")
MIRROR = str.maketrans("()[]{}<>«»", ")(][}{><»«")

BUNDLED_REGULAR = os.path.join(FONT_DIR, "arabic.ttf")
BUNDLED_BOLD = os.path.join(FONT_DIR, "arabic-bold.ttf")

SYSTEM_ARABIC_FONTS = [
    r"C:\Windows\Fonts\arial.ttf", r"C:\Windows\Fonts\tahoma.ttf",
    r"C:\Windows\Fonts\segoeui.ttf",
    "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
    "/System/Library/Fonts/Supplemental/Arial.ttf",
    "/Library/Fonts/Arial Unicode.ttf",
    "/usr/share/fonts/truetype/noto/NotoNaskhArabic-Regular.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/truetype/freefont/FreeSerif.ttf",
]
_font_cache = {}


def arabic_font(bold=False):
    """Bundled Amiri first, then any Arabic-capable system font, then any
    font the user dropped into ./fonts."""
    key = "b" if bold else "r"
    if key in _font_cache:
        return _font_cache[key]
    candidates = ([BUNDLED_BOLD] if bold else []) + [BUNDLED_REGULAR] + SYSTEM_ARABIC_FONTS
    found = next((p for p in candidates if os.path.exists(p)), None)
    if not found:
        d = FONT_DIR
        if os.path.isdir(d):
            for n in sorted(os.listdir(d)):
                if n.lower().endswith((".ttf", ".otf")):
                    found = os.path.join(d, n)
                    break
    _font_cache[key] = found
    return found


def has_arabic(s):
    return bool(RE_ARABIC.search(s or ""))


def _check_shaping():
    """Confirm at startup that Arabic shaping actually works. Without it every
    Arabic watermark, signature line and form value comes out as disconnected
    letters in reverse — and it would do so silently, which is worse than
    failing. Windows Smart App Control blocking an unsigned binary is one real
    way this happens in the wild."""
    try:
        import arabic_reshaper
        from bidi.algorithm import get_display
        probe = get_display(arabic_reshaper.reshape("اختبار"))
        if not probe or has_arabic(probe) is False:
            return False, "shaping produced no Arabic output"
        return True, ""
    except Exception as e:
        return False, f"{type(e).__name__}: {e}"


SHAPING_OK, SHAPING_ERROR = _check_shaping()
SHAPING_WARNING = (
    " ⚠ Arabic letter shaping is unavailable on this machine, so Arabic text "
    "was written unshaped and in the wrong order. "
    f"Reason: {SHAPING_ERROR}"
)


def shape_ar(text):
    """Logical Arabic -> shaped, visually ordered glyphs.
    PDF libraries do no shaping of their own, so we must do it here."""
    if not has_arabic(text) or not SHAPING_OK:
        return text
    try:
        import arabic_reshaper
        from bidi.algorithm import get_display
        return get_display(arabic_reshaper.reshape(text))
    except Exception:
        return text


def shaping_note(*texts):
    """Suffix warning callers append when they wrote Arabic without shaping."""
    if SHAPING_OK or not any(has_arabic(t) for t in texts if t):
        return ""
    return SHAPING_WARNING


def visual_to_logical(line):
    """Reverse of the above: text pulled out of a PDF in rendered order,
    turned back into readable logical order.

    Order matters — reverse first, THEN expand ligatures. A lam-alef glyph
    is one character on the page but two letters in logical order, so
    normalising before the reverse would flip that pair the wrong way
    ('الأداء' would come out as 'األداء')."""
    rev = line[::-1].translate(MIRROR)
    rev = unicodedata.normalize("NFKC", rev)
    return RE_LTR_RUN.sub(lambda m: m.group(0)[::-1], rev)


def unshape_ar(text, mode="auto"):
    """Clean up Arabic extracted from a PDF.

    mode 'auto'  — repair only lines stored as presentation forms
    mode 'force' — repair every line containing Arabic
    mode 'off'   — return the text untouched
    """
    if mode == "off" or not text:
        return text
    out = []
    for line in text.split("\n"):
        shaped = bool(RE_PRESENTATION.search(line))
        if shaped or (mode == "force" and has_arabic(line)):
            line = visual_to_logical(line)
        out.append(line)
    return "\n".join(out)


def pdf_is_arabic(path, sample_pages=3):
    """Quick look at the first few pages to decide whether this is an
    Arabic document."""
    try:
        doc = pymupdf.open(path)
        if doc.needs_pass:
            doc.authenticate("")
        txt = ""
        for i in range(min(sample_pages, doc.page_count)):
            txt += doc[i].get_text()[:3000]
        doc.close()
        if not txt.strip():
            return False
        ar = len(RE_ARABIC.findall(txt))
        return ar > 20 and ar > len(re.findall(r"[A-Za-z]", txt)) * 0.3
    except Exception:
        return False


def set_rtl(paragraph):
    """Mark a python-docx paragraph as right-to-left."""
    pPr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    pPr.append(bidi)
    jc = OxmlElement("w:jc")
    jc.set(qn("w:val"), "right")
    pPr.append(jc)
    for run in paragraph.runs:
        rPr = run._r.get_or_add_rPr()
        rtl = OxmlElement("w:rtl")
        rtl.set(qn("w:val"), "1")
        rPr.append(rtl)


# ============================================================================
# Page + file management
# ============================================================================
@app.route("/")
def index():
    return render_template("index.html")


@app.route("/api/files", methods=["GET"])
def list_files():
    return jsonify(ok=True, files=[public(f) for f in FILES.values()])


@app.route("/api/files", methods=["POST"])
def upload():
    ups = request.files.getlist("files")
    if not ups:
        return err("No file selected.")
    added = []
    for up in ups:
        name = safe_name(up.filename)
        ext = os.path.splitext(name)[1].lower()
        kind = "pdf" if ext == ".pdf" else (
            "image" if ext in (".png", ".jpg", ".jpeg", ".webp", ".bmp", ".gif") else None)
        if kind is None:
            continue
        p = os.path.join(WORK, f"{new_fid('up')}__{name}")
        up.save(p)
        added.append(register(p, name, kind, "upload"))
    if not added:
        return err("Supported files: PDF, or images (PNG / JPG) for signatures and logos.")
    return jsonify(ok=True, files=added)


@app.route("/api/files/<fid>", methods=["DELETE"])
def delete_file(fid):
    info = FILES.pop(fid, None)
    if info:
        try:
            os.remove(info["path"])
        except Exception:
            pass
    return jsonify(ok=True)


@app.route("/api/download/<fid>")
def download(fid):
    info = FILES.get(fid)
    if not info or not os.path.exists(info["path"]):
        abort(404, description="File not found.")
    return send_file(info["path"], as_attachment=True, download_name=info["name"])


@app.route("/api/preview/<fid>")
def preview(fid):
    info = FILES.get(fid)
    if not info or not os.path.exists(info["path"]):
        abort(404, description="File not found.")
    mt = mimetypes.guess_type(info["name"])[0] or "application/octet-stream"
    return send_file(info["path"], mimetype=mt)


@app.route("/api/thumb/<fid>/<int:page>")
def thumb(fid, page):
    info = FILES.get(fid)
    if not info:
        abort(404, description="File not found.")
    w = int(request.args.get("w", 190))
    try:
        if info["kind"] == "image":
            im = Image.open(info["path"])
            if im.mode in ("RGBA", "LA", "P"):
                im = im.convert("RGBA")
                bg = Image.new("RGB", im.size, (255, 255, 255))
                bg.paste(im, mask=im.split()[-1])
                im = bg
            else:
                im = im.convert("RGB")
            im.thumbnail((w, w * 3))
            buf = io.BytesIO()
            im.save(buf, "JPEG", quality=82)
            buf.seek(0)
            return send_file(buf, mimetype="image/jpeg")
        doc = pymupdf.open(info["path"])
        if doc.needs_pass:
            doc.authenticate("")
        pg = doc[page]
        zoom = w / max(pg.rect.width, 1)
        data = pg.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), alpha=False).tobytes(
            "jpeg", jpg_quality=80)
        doc.close()
        return send_file(io.BytesIO(data), mimetype="image/jpeg")
    except Exception as e:
        abort(404, description=f"Could not render a preview: {e}")


# ============================================================================
# 1) Merge
# ============================================================================
@app.route("/api/merge", methods=["POST"])
def merge():
    d = request.get_json(force=True)
    ids = d.get("ids") or []
    if len(ids) < 2:
        return err("Pick at least two PDFs to merge.")
    writer = PdfWriter()
    for fid in ids:
        info = get_file(fid, "pdf")
        for pg in reader_of(info).pages:
            writer.add_page(pg)
    return finish(writer, d.get("out") or "merged.pdf")


# ============================================================================
# 2) Organize pages (delete / rotate / reorder)
# ============================================================================
@app.route("/api/organize", methods=["POST"])
def organize():
    d = request.get_json(force=True)
    info = get_file(d.get("id"), "pdf")
    plan = d.get("plan") or []
    if not plan:
        return err("No pages left — a PDF can't be empty.")
    reader = reader_of(info)
    total = len(reader.pages)
    writer = PdfWriter()
    for item in plan:
        src = int(item.get("src", 0))
        if not (0 <= src < total):
            continue
        pg = reader.pages[src]
        rot = int(item.get("rotate", 0)) % 360
        if rot:
            pg.rotate(rot)
        writer.add_page(pg)
    if not len(writer.pages):
        return err("No valid pages to save.")
    return finish(writer, d.get("out") or f"{stem(info['name'])}_organized.pdf")


# ============================================================================
# 3) Split
# ============================================================================
@app.route("/api/split", methods=["POST"])
def split():
    d = request.get_json(force=True)
    info = get_file(d.get("id"), "pdf")
    mode = d.get("mode", "each")
    reader = reader_of(info)
    total = len(reader.pages)
    base = stem(info["name"])
    chunks = []

    if mode == "each":
        chunks = [(f"page_{i+1}", [i]) for i in range(total)]
    elif mode == "every":
        n = max(1, int(d.get("every", 2)))
        for s in range(0, total, n):
            idx = list(range(s, min(s + n, total)))
            chunks.append((f"{idx[0]+1}-{idx[-1]+1}", idx))
    elif mode == "ranges":
        parts = [p.strip() for p in str(d.get("ranges", "")).replace("،", ",").split(",") if p.strip()]
        if not parts:
            return err("Enter ranges such as: 1-3, 4-7, 10")
        for part in parts:
            idx = parse_pages(part, total)
            if idx:
                chunks.append((part.replace(" ", ""), idx))
    elif mode == "extract":
        idx = parse_pages(d.get("ranges", "all"), total)
        writer = PdfWriter()
        for i in idx:
            writer.add_page(reader.pages[i])
        return finish(writer, f"{base}_extracted.pdf")
    else:
        return err("Unknown split mode.")

    if not chunks:
        return err("Nothing to split.")

    if len(chunks) == 1:
        writer = PdfWriter()
        for i in chunks[0][1]:
            writer.add_page(reader.pages[i])
        return finish(writer, f"{base}_{chunks[0][0]}.pdf")

    zp = out_path(f"{base}_split.zip")
    with zipfile.ZipFile(zp, "w", zipfile.ZIP_DEFLATED) as z:
        for label, idx in chunks:
            w = PdfWriter()
            for i in idx:
                w.add_page(reader.pages[i])
            buf = io.BytesIO()
            w.write(buf)
            z.writestr(f"{base}_{label}.pdf", buf.getvalue())
    return jsonify(ok=True, file=register(zp, f"{base}_split.zip", "zip", "output"),
                   note=f"{len(chunks)} files created inside a zip archive.")


# ============================================================================
# 4) Extract / convert
# ============================================================================
@app.route("/api/extract", methods=["POST"])
def extract():
    d = request.get_json(force=True)
    info = get_file(d.get("id"), "pdf")
    kind = d.get("kind", "text")
    ar_mode = d.get("arabic_fix", "auto")   # auto | force | off
    base = stem(info["name"])
    reader = reader_of(info)
    total = len(reader.pages)
    idx = parse_pages(d.get("pages", "all"), total)

    # ---------------- plain text ----------------
    if kind in ("text", "text_layout"):
        chunks, repaired = [], 0
        with pdfplumber.open(info["path"]) as pdf:
            for i in idx:
                raw = pdf.pages[i].extract_text(layout=(kind == "text_layout")) or ""
                fixed = unshape_ar(raw, ar_mode)
                if fixed != raw:
                    repaired += 1
                chunks.append(f"===== Page {i+1} =====\n{fixed}")
        body = "\n\n".join(chunks)
        p = out_path(f"{base}.txt")
        with open(p, "w", encoding="utf-8") as fh:
            fh.write(body)
        note = "Text extracted."
        if repaired:
            note = f"Arabic letter order repaired on {repaired} page(s)."
        return jsonify(ok=True, file=register(p, f"{base}.txt", "txt", "output"),
                       note=note, text_preview=body[:4000])

    # ---------------- Word ----------------
    if kind == "docx":
        p = out_path(f"{base}.docx")
        want = d.get("docx_mode", "auto")     # auto | layout | text
        arabic_doc = pdf_is_arabic(info["path"])
        use_text = (want == "text") or (want == "auto" and arabic_doc)
        fallback = ""

        if not use_text:
            try:
                from pdf2docx import Converter
                cv = Converter(info["path"])
                cv.convert(p, start=min(idx), end=max(idx) + 1)
                cv.close()
                return jsonify(ok=True, file=register(p, f"{base}.docx", "docx", "output"),
                               note="Converted with layout preserved (pdf2docx).")
            except Exception as e:
                use_text = True
                fallback = f" Layout conversion failed ({e}), fell back to text mode."

        doc = DocxDocument()
        doc.styles["Normal"].font.size = Pt(12)
        doc.styles["Normal"].font.name = "Arial"
        repaired = 0
        with pdfplumber.open(info["path"]) as pdf:
            for i in idx:
                raw = pdf.pages[i].extract_text() or ""
                fixed = unshape_ar(raw, ar_mode)
                if fixed != raw:
                    repaired += 1
                for line in fixed.split("\n"):
                    par = doc.add_paragraph(line)
                    if has_arabic(line):
                        set_rtl(par)
                if i != idx[-1]:
                    doc.add_page_break()
        doc.save(p)
        note = "Converted in text mode"
        if arabic_doc:
            note += " (Arabic document — paragraphs set right-to-left)"
        if repaired:
            note += f", letter order repaired on {repaired} page(s)"
        return jsonify(ok=True, file=register(p, f"{base}.docx", "docx", "output"),
                       note=note + "." + fallback)

    # ---------------- Excel ----------------
    if kind in ("tables", "xlsx"):
        wb = Workbook()
        wb.remove(wb.active)
        found, rtl_sheets = 0, 0
        with pdfplumber.open(info["path"]) as pdf:
            for i in idx:
                for j, table in enumerate(pdf.pages[i].extract_tables() or []):
                    if not table:
                        continue
                    found += 1
                    ws = wb.create_sheet(f"p{i+1}_t{j+1}"[:31])
                    sheet_ar = False
                    for row in table:
                        cells = []
                        for c in row:
                            v = "" if c is None else unshape_ar(str(c), ar_mode)
                            if has_arabic(v):
                                sheet_ar = True
                            cells.append(v)
                        ws.append(cells)
                    if sheet_ar:
                        ws.sheet_view.rightToLeft = True
                        rtl_sheets += 1
                        for row in ws.iter_rows():
                            for cell in row:
                                cell.alignment = Alignment(horizontal="right",
                                                           readingOrder=2)
                    for col in ws.columns:
                        width = min(60, max((len(str(c.value or "")) for c in col), default=8) + 2)
                        ws.column_dimensions[col[0].column_letter].width = width
        if not found:
            return err("No tables found on the selected pages.")
        p = out_path(f"{base}_tables.xlsx")
        wb.save(p)
        note = f"{found} table(s) extracted."
        if rtl_sheets:
            note += f" {rtl_sheets} sheet(s) set right-to-left."
        return jsonify(ok=True, file=register(p, f"{base}_tables.xlsx", "xlsx", "output"),
                       note=note)

    # ---------------- embedded images ----------------
    if kind == "images":
        doc = pymupdf.open(info["path"])
        if doc.needs_pass:
            doc.authenticate("")
        zp = out_path(f"{base}_images.zip")
        n = 0
        with zipfile.ZipFile(zp, "w", zipfile.ZIP_DEFLATED) as z:
            for i in idx:
                for k, img in enumerate(doc[i].get_images(full=True)):
                    try:
                        d2 = doc.extract_image(img[0])
                        z.writestr(f"page{i+1}_img{k+1}.{d2['ext']}", d2["image"])
                        n += 1
                    except Exception:
                        pass
        doc.close()
        if not n:
            os.remove(zp)
            return err("No embedded images on the selected pages.")
        return jsonify(ok=True, file=register(zp, f"{base}_images.zip", "zip", "output"),
                       note=f"{n} image(s) extracted.")

    # ---------------- pages as PNG ----------------
    if kind == "pages_png":
        doc = pymupdf.open(info["path"])
        if doc.needs_pass:
            doc.authenticate("")
        dpi = int(d.get("dpi", 150))
        zp = out_path(f"{base}_pages_png.zip")
        with zipfile.ZipFile(zp, "w", zipfile.ZIP_DEFLATED) as z:
            for i in idx:
                z.writestr(f"{base}_page{i+1}.png", doc[i].get_pixmap(dpi=dpi).tobytes("png"))
        doc.close()
        return jsonify(ok=True, file=register(zp, f"{base}_pages_png.zip", "zip", "output"),
                       note=f"{len(idx)} page(s) exported at {dpi} DPI.")

    return err("Unknown extraction type.")


# ============================================================================
# 5) Compress
# ============================================================================
# Three levels, and the first one is the point of the whole section: a file can
# usually be made smaller without touching a single pixel. PDFs accumulate
# slack — streams stored at a low compression setting, the same logo embedded
# once per page, objects nothing points at any more — and clearing that out is
# arithmetic, not a quality trade-off. Only the other two levels resample
# images, and the result is checked against the original before it is returned.

LEVELS = {
    # name: (dpi above which an image is resampled, target dpi, JPEG quality)
    "lossless": None,
    "balanced": (200, 150, 82),
    "maximum":  (120, 96, 62),
}


def _open_pdf(path):
    doc = pymupdf.open(path)
    if doc.needs_pass:
        doc.authenticate("")
    return doc


def _save_compact(doc, dest):
    """Save with everything the writer can do that never changes appearance."""
    opts = dict(garbage=4, deflate=True, deflate_images=True,
                deflate_fonts=True, clean=True)
    for extra in ({"use_objstms": 1}, {}):
        try:
            doc.save(dest, **opts, **extra)
            return
        except TypeError:
            continue          # older PyMuPDF without that keyword
    doc.save(dest)


def dedupe_images(src, dest):
    """Point every copy of the same image at one object.

    A logo repeated on forty pages is usually forty separate image streams.
    They are byte-identical, so the page can reference the first one and the
    rest become unreachable — the garbage pass then drops them. Returns the
    number of copies removed."""
    removed = 0
    try:
        with pikepdf.open(src) as pdf:
            first = {}
            for page in pdf.pages:
                try:
                    xobjects = page["/Resources"]["/XObject"]
                except Exception:
                    continue
                for name in list(xobjects.keys()):
                    obj = xobjects[name]
                    try:
                        if obj.get("/Subtype") != pikepdf.Name("/Image"):
                            continue
                        digest = hashlib.sha1(obj.read_raw_bytes()).digest()
                    except Exception:
                        continue
                    if digest in first:
                        if obj.objgen != first[digest].objgen:
                            xobjects[name] = first[digest]
                            removed += 1
                    else:
                        first[digest] = obj
            if removed:
                pdf.save(dest, compress_streams=True, recompress_flate=True,
                         object_stream_mode=pikepdf.ObjectStreamMode.generate)
                return removed
    except Exception:
        return 0
    return 0


def image_report(doc):
    """Every placed image with the resolution it is actually displayed at.

    The stored pixel count only matters relative to the space it occupies on
    the page: a 4000px-wide photo in a 2-inch box is 2000 DPI, and nothing can
    show that. This is what decides whether resampling would help."""
    items, seen = [], set()
    for page in doc:
        try:
            placed = page.get_images(full=True)
        except Exception:
            continue
        for entry in placed:
            xref = entry[0]
            try:
                rects = page.get_image_rects(xref)
            except Exception:
                rects = []
            width, height = entry[2], entry[3]
            dpi = 0
            if rects:
                box = rects[0]
                if box.width > 1 and box.height > 1:
                    dpi = max(width * 72.0 / box.width, height * 72.0 / box.height)
            try:
                size = len(doc.xref_stream_raw(xref) or b"")
            except Exception:
                size = 0
            key = (xref, round(dpi))
            if key in seen:
                continue
            seen.add(key)
            items.append({"page": page.number + 1, "w": width, "h": height,
                          "dpi": round(dpi), "bytes": size})
    items.sort(key=lambda i: i["bytes"], reverse=True)
    return items


def composition(path):
    """How the bytes are spent: images, fonts, everything else."""
    buckets = {"images": 0, "fonts": 0, "other": 0}
    try:
        with pikepdf.open(path) as pdf:
            for obj in pdf.objects:
                try:
                    if not isinstance(obj, pikepdf.Stream):
                        continue
                    size = len(obj.read_raw_bytes())
                except Exception:
                    continue
                subtype = str(obj.get("/Subtype", ""))
                if subtype == "/Image":
                    buckets["images"] += size
                elif "/Length1" in obj or subtype in ("/Type1C", "/CIDFontType0C",
                                                      "/OpenType"):
                    buckets["fonts"] += size
                else:
                    buckets["other"] += size
    except Exception:
        pass
    return buckets


def visual_delta(before, after, pages=3, dpi=110):
    """How different the two files look, 0 meaning not at all.

    Returned as the mean per-pixel difference on a 0-100 scale, so a lossless
    run can be shown to be exactly that rather than merely claimed to be."""
    try:
        from PIL import Image, ImageChops, ImageStat
    except Exception:
        return None
    try:
        a, b = _open_pdf(before), _open_pdf(after)
    except Exception:
        return None
    worst = 0.0
    try:
        if a.page_count != b.page_count:
            return None
        for n in range(min(pages, a.page_count)):
            pa = a[n].get_pixmap(dpi=dpi)
            pb = b[n].get_pixmap(dpi=dpi)
            if (pa.width, pa.height) != (pb.width, pb.height):
                return None
            ia = Image.frombytes("RGB", (pa.width, pa.height), pa.samples)
            ib = Image.frombytes("RGB", (pb.width, pb.height), pb.samples)
            stat = ImageStat.Stat(ImageChops.difference(ia, ib))
            worst = max(worst, sum(stat.mean) / len(stat.mean) * 100.0 / 255.0)
    except Exception:
        return None
    finally:
        try:
            a.close(); b.close()
        except Exception:
            pass
    return round(worst, 3)


@app.route("/api/analyze", methods=["POST"])
def analyze():
    """What is taking up the room, and what would actually help."""
    d = request.get_json(force=True)
    info = get_file(d.get("id"), "pdf")
    total = os.path.getsize(info["path"])
    parts = composition(info["path"])
    try:
        doc = _open_pdf(info["path"])
    except Exception as e:
        return err(f"Could not read that file ({e}).")
    try:
        images = image_report(doc)
        page_count = doc.page_count
    finally:
        doc.close()

    oversized = [i for i in images if i["dpi"] > 200]
    oversized_bytes = sum(i["bytes"] for i in oversized)
    share = (parts["images"] / total * 100) if total else 0

    if oversized_bytes > total * 0.25:
        advice = (f"{human(oversized_bytes)} of this file is images stored at more "
                  f"resolution than the page can show. <b>Smaller</b> will cut most "
                  f"of that, and on screen or in print it will look the same.")
    elif share > 40:
        advice = ("Most of this file is images, but they are already at a sensible "
                  "resolution. <b>Smaller</b> still helps; <b>Smallest</b> will "
                  "visibly soften them.")
    else:
        advice = ("This file is mostly text and vector drawing, so resampling images "
                  "has little to work with. <b>Keep quality exactly</b> is the right "
                  "choice — anything more trades quality for very little.")

    return jsonify(ok=True, total=total, total_h=human(total), pages=page_count,
                   parts={k: {"bytes": v, "h": human(v),
                              "pct": round(v / total * 100, 1) if total else 0}
                          for k, v in parts.items()},
                   images=images[:6], image_count=len(images),
                   oversized=len(oversized), oversized_h=human(oversized_bytes),
                   advice=advice)


@app.route("/api/compress", methods=["POST"])
def compress():
    d = request.get_json(force=True)
    info = get_file(d.get("id"), "pdf")
    level = d.get("level", "balanced")
    if level == "light":            # the name used before this was rewritten
        level = "lossless"
    if level in ("medium", "strong"):
        level = {"medium": "balanced", "strong": "maximum"}[level]
    if level not in LEVELS:
        return err("Unknown compression level.")

    base = stem(info["name"])
    src = info["path"]
    before = os.path.getsize(src)
    work = os.path.join(WORK, f"{new_fid('tmp')}.pdf")
    shutil.copy(src, work)
    notes = []

    # Shared by every level: identical images collapsed into one.
    staged = os.path.join(WORK, f"{new_fid('tmp')}.pdf")
    copies = dedupe_images(work, staged)
    if copies:
        os.replace(staged, work)
        notes.append(f"{copies} repeated image{'s' if copies > 1 else ''} stored once")
    elif os.path.exists(staged):
        try:
            os.remove(staged)
        except OSError:
            pass

    settings = LEVELS[level]
    if settings:
        threshold, target_dpi, quality = settings
        staged = os.path.join(WORK, f"{new_fid('tmp')}.pdf")
        try:
            doc = _open_pdf(work)
            over = sum(1 for i in image_report(doc) if i["dpi"] > threshold)
            doc.rewrite_images(dpi_threshold=threshold, dpi_target=target_dpi,
                               quality=quality, lossy=True, lossless=True)
            _save_compact(doc, staged)
            doc.close()
            os.replace(staged, work)
            notes.append(
                f"{over} image{'s' if over != 1 else ''} above {threshold} DPI resampled "
                f"to {target_dpi} DPI" if over else
                f"no images above {threshold} DPI, so none needed resampling")
        except Exception as e:
            notes.append(f"image pass skipped ({e})")
            if os.path.exists(staged):
                try:
                    os.remove(staged)
                except OSError:
                    pass
    else:
        # Nothing is re-encoded. Fonts are deliberately left whole: subsetting
        # them would save a little, but it strips the character map, and the
        # Edit page tab needs that to set new text in the document's own font.
        staged = os.path.join(WORK, f"{new_fid('tmp')}.pdf")
        try:
            doc = _open_pdf(work)
            _save_compact(doc, staged)
            doc.close()
            os.replace(staged, work)
        except Exception:
            if os.path.exists(staged):
                try:
                    os.remove(staged)
                except OSError:
                    pass

    out = out_path(f"{base}_compressed.pdf")
    try:
        with pikepdf.open(work) as pdf:
            pdf.save(out, compress_streams=True, recompress_flate=True,
                     object_stream_mode=pikepdf.ObjectStreamMode.generate,
                     linearize=False)
    except Exception:
        shutil.copy(work, out)
    try:
        os.remove(work)
    except OSError:
        pass

    after = os.path.getsize(out)
    kept_original = False
    if after >= before:
        shutil.copy(src, out)
        after = os.path.getsize(out)
        kept_original = True

    # Say what happened to the pages, rather than asking to be trusted.
    delta = visual_delta(src, out) if not kept_original else 0.0
    if delta is None:
        check = "the pages could not be compared automatically"
    elif delta == 0:
        check = "checked against the original: the pages are pixel-for-pixel identical"
    elif delta < 0.5:
        check = f"checked against the original: {delta}% average pixel difference — not visible"
    elif delta < 2:
        check = f"checked against the original: {delta}% average pixel difference — very slight softening"
    else:
        check = f"checked against the original: {delta}% average pixel difference — visibly softer"

    if kept_original:
        note = (f"{human(before)} — already as small as it goes at this level, "
                f"so the original was kept untouched.")
    else:
        pct = (1 - after / before) * 100 if before else 0
        note = f"{human(before)} → {human(after)}, {pct:.1f}% smaller. " + \
               (" · ".join(notes) + ". " if notes else "") + check + "."

    return jsonify(ok=True, file=register(out, f"{base}_compressed.pdf", "pdf", "output"),
                   saved_pct=round((1 - after / before) * 100, 1) if before else 0,
                   delta=delta, note=note)


# ============================================================================
# 6) Security
# ============================================================================
@app.route("/api/protect", methods=["POST"])
def protect():
    d = request.get_json(force=True)
    info = get_file(d.get("id"), "pdf")
    user_pw = (d.get("user_pw") or "").strip()
    owner_pw = (d.get("owner_pw") or "").strip() or user_pw
    if not user_pw:
        return err("Enter an open password.")
    base = stem(info["name"])
    p = out_path(f"{base}_protected.pdf")
    perms = pikepdf.Permissions(
        extract=bool(d.get("allow_copy", True)),
        modify_annotation=bool(d.get("allow_annotate", True)),
        modify_assembly=False,
        modify_form=bool(d.get("allow_form", True)),
        modify_other=bool(d.get("allow_edit", False)),
        print_lowres=bool(d.get("allow_print", True)),
        print_highres=bool(d.get("allow_print", True)),
    )
    with pikepdf.open(info["path"]) as pdf:
        pdf.save(p, encryption=pikepdf.Encryption(user=user_pw, owner=owner_pw, R=6, allow=perms))
    out = register(p, f"{base}_protected.pdf", "pdf", "output")
    out["encrypted"] = True
    return jsonify(ok=True, file=out, note="Encrypted with AES-256.")


@app.route("/api/unlock", methods=["POST"])
def unlock():
    d = request.get_json(force=True)
    info = FILES.get(d.get("id"))
    if not info:
        return err("File not found.")
    base = stem(info["name"])
    p = out_path(f"{base}_unlocked.pdf")
    try:
        with pikepdf.open(info["path"], password=d.get("pw") or "") as pdf:
            pdf.save(p)
    except pikepdf.PasswordError:
        return err("Wrong password.")
    except Exception as e:
        return err(f"Could not open the file: {e}")
    return jsonify(ok=True, file=register(p, f"{base}_unlocked.pdf", "pdf", "output"),
                   note="Password removed — the new file opens without one.")


# ============================================================================
# 7) Watermark
# ============================================================================
def register_rl_font(bold=False):
    """Register the Arabic TTF with ReportLab and return its name."""
    fp = arabic_font(bold)
    if not fp:
        return None
    name = "PDFStudioAR-B" if bold else "PDFStudioAR"
    try:
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        if name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(name, fp))
        return name
    except Exception:
        return None


def make_overlay(w, h, opts):
    buf = io.BytesIO()
    c = rl_canvas.Canvas(buf, pagesize=(w, h))
    c.saveState()
    try:
        c.setFillAlpha(float(opts.get("opacity", 0.18)))
        c.setStrokeAlpha(float(opts.get("opacity", 0.18)))
    except Exception:
        pass

    if opts.get("mode") == "image" and opts.get("_img"):
        img = ImageReader(opts["_img"])
        iw, ih = img.getSize()
        scale = float(opts.get("scale", 0.4)) * w / iw
        dw, dh = iw * scale, ih * scale
        c.translate(w * float(opts.get("x", 0.5)), h * float(opts.get("y", 0.5)))
        c.rotate(float(opts.get("angle", 0)))
        c.drawImage(img, -dw / 2, -dh / 2, dw, dh, mask="auto")
    else:
        raw = opts.get("text") or "WATERMARK"
        text = shape_ar(raw)
        size = float(opts.get("size", 48))
        c.setFillColor(HexColor(hex_norm(opts.get("color"), "#FF0000")))
        font = "Helvetica-Bold"
        if has_arabic(raw):
            font = register_rl_font(bold=True) or register_rl_font() or font
        c.setFont(font, size)
        if opts.get("tile"):
            step_x = max(c.stringWidth(text, font, size) + size * 2, size * 4)
            step_y = size * 4
            c.saveState()
            c.translate(w / 2, h / 2)
            c.rotate(float(opts.get("angle", 45)))
            c.translate(-w / 2, -h / 2)
            y = -h
            while y < h * 2:
                x = -w
                while x < w * 2:
                    c.drawString(x, y, text)
                    x += step_x
                y += step_y
            c.restoreState()
        else:
            c.translate(w * float(opts.get("x", 0.5)), h * float(opts.get("y", 0.5)))
            c.rotate(float(opts.get("angle", 45)))
            c.drawCentredString(0, -size * 0.35, text)
    c.restoreState()
    c.save()
    buf.seek(0)
    return PdfReader(buf).pages[0]


@app.route("/api/watermark", methods=["POST"])
def watermark():
    d = request.get_json(force=True)
    info = get_file(d.get("id"), "pdf")
    opts = dict(d)
    if d.get("mode") == "image":
        opts["_img"] = get_file(d.get("image_id"), "image")["path"]
    elif not (d.get("text") or "").strip():
        return err("Enter the watermark text.")

    reader = reader_of(info)
    targets = set(parse_pages(d.get("pages", "all"), len(reader.pages)))
    writer = PdfWriter()
    for i, page in enumerate(reader.pages):
        if i in targets:
            w, h = float(page.mediabox.width), float(page.mediabox.height)
            over = make_overlay(w, h, opts)   # regenerated per page: merging consumes it
            if d.get("behind"):
                over.merge_page(page)
                writer.add_page(over)
                continue
            page.merge_page(over)
        writer.add_page(page)
    return finish(writer, f"{stem(info['name'])}_watermarked.pdf",
                  note=(shaping_note(d.get("text")) or None))


# ============================================================================
# 8) Signature / stamp / free text
# ============================================================================
@app.route("/api/stamp", methods=["POST"])
def stamp():
    d = request.get_json(force=True)
    info = get_file(d.get("id"), "pdf")
    img = get_file(d.get("image_id"), "image")
    doc = pymupdf.open(info["path"])
    if doc.needs_pass:
        doc.authenticate("")
    page_no = int(d.get("page", 0))
    if not (0 <= page_no < doc.page_count):
        doc.close()
        return err("Page number out of range.")

    pages = parse_pages(d.get("pages") or str(page_no + 1), doc.page_count)
    with Image.open(img["path"]) as im:
        iw, ih = im.size
    ratio = ih / iw if iw else 1

    for pn in pages:
        pg = doc[pn]
        pw, ph = pg.rect.width, pg.rect.height
        wpx = float(d.get("width", 0.22)) * pw
        hpx = wpx * ratio
        cx = float(d.get("x", 0.7)) * pw
        cy = (1 - float(d.get("y", 0.15))) * ph   # y measured from the bottom
        pg.insert_image(pymupdf.Rect(cx - wpx / 2, cy - hpx / 2, cx + wpx / 2, cy + hpx / 2),
                        filename=img["path"], overlay=True, keep_proportion=True,
                        rotate=int(d.get("rotate", 0)) // 90 * 90 % 360)

    p = out_path(f"{stem(info['name'])}_signed.pdf")
    doc.save(p, garbage=3, deflate=True)
    doc.close()
    return jsonify(ok=True, file=register(p, f"{stem(info['name'])}_signed.pdf", "pdf", "output"),
                   note=f"Image placed on {len(pages)} page(s).")


def draw_text(pg, txt, x_frac, y_frac, size, rgb, align=1):
    """Draw a line of text (Arabic-aware) centred on a point given as
    fractions of the page. y is measured from the bottom."""
    shaped = shape_ar(txt)
    fontfile = arabic_font() if has_arabic(txt) else None
    x = x_frac * pg.rect.width
    y = (1 - y_frac) * pg.rect.height
    rect = pymupdf.Rect(x - pg.rect.width * 0.47, y - size * 2.2,
                        x + pg.rect.width * 0.47, y + size * 2.2)
    fontname = "arb" if fontfile else "helv"
    s = size
    while s >= 4.5:
        try:
            if pg.insert_textbox(rect, shaped, fontsize=s, color=rgb,
                                 fontname=fontname, fontfile=fontfile, align=align) >= 0:
                return
        except Exception:
            break
        s -= 0.5
    try:
        pg.insert_text((rect.x0, y), shaped, fontsize=size, color=rgb,
                       fontname=fontname, fontfile=fontfile)
    except Exception:
        pg.insert_text((rect.x0, y), shaped, fontsize=size, color=rgb, fontname="helv")


@app.route("/api/text", methods=["POST"])
def add_text():
    d = request.get_json(force=True)
    info = get_file(d.get("id"), "pdf")
    txt = (d.get("text") or "").strip()
    if not txt:
        return err("Enter the text to add.")
    doc = pymupdf.open(info["path"])
    if doc.needs_pass:
        doc.authenticate("")
    pages = parse_pages(d.get("pages") or "1", doc.page_count)
    size = float(d.get("size", 14))
    rgb = hex_rgb(d.get("color"), (0, 0, 0))
    for pn in pages:
        draw_text(doc[pn], txt, float(d.get("x", 0.5)), float(d.get("y", 0.1)), size, rgb)
    p = out_path(f"{stem(info['name'])}_text.pdf")
    doc.save(p, garbage=3, deflate=True)
    doc.close()
    note = f"Text added to {len(pages)} page(s)."
    if has_arabic(txt) and arabic_font() == BUNDLED_REGULAR:
        note += " Rendered with the bundled Amiri font."
    note += shaping_note(txt)
    return jsonify(ok=True, file=register(p, f"{stem(info['name'])}_text.pdf", "pdf", "output"),
                   note=note)


# ============================================================================
# 9) PDF forms
# ============================================================================
FIELD_TYPES = {"/Tx": "Text", "/Btn": "Button / Checkbox", "/Ch": "Choice", "/Sig": "Signature"}


@app.route("/api/form/<fid>")
def form_fields(fid):
    info = get_file(fid, "pdf")
    try:
        fields = reader_of(info).get_fields() or {}
    except Exception as e:
        return err(f"Could not read the form fields: {e}")
    out = []
    for name, f in fields.items():
        states = None
        try:
            kids = f.get("/_States_")
            if kids:
                states = [str(s) for s in kids]
        except Exception:
            pass
        out.append({
            "name": str(name),
            "type": FIELD_TYPES.get(str(f.get("/FT", "")), str(f.get("/FT", "")) or "Unknown"),
            "value": "" if f.get("/V") is None else str(f.get("/V")),
            "states": states,
        })
    out.sort(key=lambda x: x["name"])
    return jsonify(ok=True, fields=out, count=len(out))


def fit_textbox(pg, box, text, fontfile, align, start_size):
    """insert_textbox refuses to draw when the text doesn't fit, so step the
    size down until it does. Arabic faces like Amiri have a tall line height,
    which makes this necessary inside short form fields."""
    fontname = "arb" if fontfile else "helv"
    size = float(start_size)
    while size >= 4.5:
        rc = pg.insert_textbox(box, text, fontsize=size, fontname=fontname,
                               fontfile=fontfile, align=align, color=(0, 0, 0))
        if rc >= 0:
            return True
        size -= 0.5
    try:   # last resort: a single unwrapped baseline
        pg.insert_text((box.x0, box.y1 - 3), text, fontsize=7, fontname=fontname,
                       fontfile=fontfile, color=(0, 0, 0))
        return True
    except Exception:
        return False


def flatten_form_arabic(src_path, values, dst_path):
    """Burn field values into the page with the Arabic font, then remove the
    widgets. Needed because PDF form fields use built-in fonts that have no
    Arabic glyphs — filling them normally produces blanks or boxes."""
    doc = pymupdf.open(src_path)
    if doc.needs_pass:
        doc.authenticate("")
    fontfile = arabic_font()
    drawn = 0
    for pg in doc:
        widgets = list(pg.widgets() or [])
        for wdg in widgets:
            val = values.get(wdg.field_name)
            if val is None:
                val = wdg.field_value or ""
            val = str(val)
            if val.strip():
                r = wdg.rect
                size = min(max(r.height * 0.60, 7), 15)
                box = pymupdf.Rect(r.x0 + 2, r.y0 - 2, r.x1 - 2, r.y1 + 2)
                ar = has_arabic(val)
                if fit_textbox(pg, box, shape_ar(val), fontfile if ar else None,
                               2 if ar else 0, size):   # right-align Arabic
                    drawn += 1
        for wdg in widgets:
            try:
                pg.delete_widget(wdg)
            except Exception:
                pass
    doc.save(dst_path, garbage=3, deflate=True)
    doc.close()
    return drawn


@app.route("/api/form/fill", methods=["POST"])
def form_fill():
    d = request.get_json(force=True)
    info = get_file(d.get("id"), "pdf")
    values = {str(k): str(v) for k, v in (d.get("values") or {}).items()}
    if not values:
        return err("No values entered.")
    base = stem(info["name"])
    any_arabic = any(has_arabic(v) for v in values.values())
    flatten = bool(d.get("flatten")) or any_arabic

    if flatten:
        p = out_path(f"{base}_filled_flat.pdf")
        drawn = flatten_form_arabic(info["path"], values, p)
        note = f"{drawn} field(s) filled and flattened — the values are now part of the page."
        if any_arabic:
            note += (" Arabic values are always flattened: PDF form fields use built-in "
                     "fonts with no Arabic glyphs, so an interactive field would show blanks.")
        note += shaping_note(*values.values())
        return jsonify(ok=True, file=register(p, f"{base}_filled_flat.pdf", "pdf", "output"),
                       note=note)

    reader = reader_of(info)
    writer = PdfWriter()
    writer.append(reader)
    try:
        writer.set_need_appearances_writer(True)
    except Exception:
        pass
    for page in writer.pages:
        try:
            writer.update_page_form_field_values(page, values, auto_regenerate=False)
        except TypeError:
            writer.update_page_form_field_values(page, values)
        except Exception:
            pass
    p = out_path(f"{base}_filled.pdf")
    with open(p, "wb") as fh:
        writer.write(fh)
    return jsonify(ok=True, file=register(p, f"{base}_filled.pdf", "pdf", "output"),
                   note="Fields filled — the form stays editable.")


# ============================================================================
# 10) Editing the page: erase, retype, replace a logo
# ============================================================================
import editing


def open_doc(info):
    doc = pymupdf.open(info["path"])
    if doc.needs_pass:
        doc.authenticate("")
    return doc


def frac_rect(page, r):
    """A rectangle given as fractions of the page -> points."""
    w, h = page.rect.width, page.rect.height
    return pymupdf.Rect(float(r["x0"]) * w, float(r["y0"]) * h,
                        float(r["x1"]) * w, float(r["y1"]) * h)


@app.route("/api/inspect", methods=["POST"])
def inspect_region():
    """What is under this rectangle — so the interface can show which font,
    size and colour a replacement would be set in before anything is changed."""
    d = request.get_json(force=True)
    info = get_file(d.get("id"), "pdf")
    doc = open_doc(info)
    page_no = int(d.get("page", 0))
    if not (0 <= page_no < doc.page_count):
        doc.close()
        return err("Page number out of range.")
    page = doc[page_no]
    rect = frac_rect(page, d.get("rect") or {"x0": 0, "y0": 0, "x1": 1, "y1": 1})
    span = editing.span_at(page, rect)
    sample = d.get("text") or (span or {}).get("text") or "Sample"
    result = {"ok": True, "found": bool(span)}
    if span:
        choice = editing.resolve_font(doc, page, span, sample,
                                      arabic_font(), has_arabic(sample))
        cleanup = choice.pop("cleanup", None)
        if cleanup:
            try:
                os.remove(cleanup)
            except OSError:
                pass
        result.update({
            "font": span["raw_name"], "family": span["family"],
            "size": span["size"], "color": span["color"],
            "bold": span["bold"], "italic": span["italic"],
            "text": span["text"], "match": choice["match"], "used": choice["used"],
        })
    doc.close()
    return jsonify(**result)


def apply_ops(doc, ops, out_notes):
    """Erase everything first, then draw — redaction would otherwise wipe the
    new content along with the old."""
    by_page = {}
    for op in ops:
        by_page.setdefault(int(op.get("page", 0)), []).append(op)

    for page_no, page_ops in by_page.items():
        if not (0 <= page_no < doc.page_count):
            continue
        page = doc[page_no]

        # Remember what each region looked like before it is erased.
        for op in page_ops:
            rect = frac_rect(page, op["rect"])
            op["_rect"] = rect
            op["_span"] = editing.span_at(page, rect) if op.get("action") == "text" else None
            page.add_redact_annot(rect)

        page.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_REMOVE,
                              graphics=pymupdf.PDF_REDACT_LINE_ART_REMOVE_IF_COVERED,
                              text=pymupdf.PDF_REDACT_TEXT_REMOVE)

        for op in page_ops:
            action = op.get("action", "erase")
            rect = op["_rect"]

            if action == "text":
                text = (op.get("text") or "").strip()
                if not text:
                    continue
                span = op["_span"] or {}
                auto = op.get("auto_font", True)
                size = float(op.get("size") or (span.get("size") if auto else 0) or 12)
                rgb = span.get("rgb", (0, 0, 0)) if auto and span else hex_rgb(op.get("color"))
                arabic = has_arabic(text)
                # Sit on the original line rather than at the top of whatever
                # box was dragged, so the replacement lines up with its
                # neighbours instead of floating above them.
                if auto and span.get("bbox"):
                    sb = span["bbox"]
                    rect = pymupdf.Rect(rect.x0, min(sb[1], rect.y1 - 1),
                                        rect.x1, max(sb[3], rect.y0 + 1))
                    op["_rect"] = rect
                choice = editing.resolve_font(doc, page, span if auto else None,
                                              shape_ar(text), arabic_font(), arabic)
                cleanup = choice.pop("cleanup", None)
                if auto and span.get("origin") and "\n" not in text:
                    # One line replacing one line: put it on the original
                    # baseline, starting where the original started, rather
                    # than inside the (hand-drawn, slightly off) box.
                    room = max(rect.x1, span["bbox"][2]) - min(rect.x0, span["bbox"][0])
                    placed = editing.write_on_line(page, span, shape_ar(text), choice,
                                                   size, rgb, room, rtl=arabic)
                    used = placed[0] if placed else None
                else:
                    used = editing.write_into(page, rect, shape_ar(text), choice, size, rgb,
                                              align=int(op.get("align", 0)))
                if cleanup:
                    try:
                        os.remove(cleanup)
                    except OSError:
                        pass
                if used is None:
                    out_notes.append("some text would not fit its box and was skipped")
                else:
                    note = choice["used"]
                    if abs(used - size) > 0.4:
                        note += f", shrunk to {used:g}pt to fit"
                    out_notes.append(note)

            elif action == "image":
                img = get_file(op.get("image_id"), "image")
                try:
                    page.insert_image(rect, filename=img["path"], keep_proportion=True,
                                      overlay=True)
                except Exception as e:
                    out_notes.append(f"could not place the image: {e}")


@app.route("/api/edit", methods=["POST"])
def edit_page():
    d = request.get_json(force=True)
    info = get_file(d.get("id"), "pdf")
    ops = d.get("ops") or []
    if not ops:
        return err("Nothing selected to change.")
    doc = open_doc(info)
    notes = []
    try:
        apply_ops(doc, ops, notes)
    except Exception as e:
        doc.close()
        return err(f"Could not apply the changes: {e}")
    p = out_path(f"{stem(info['name'])}_edited.pdf")
    doc.save(p, garbage=3, deflate=True)
    doc.close()
    kinds = [o.get("action", "erase") for o in ops]
    summary = f"{len(ops)} region(s) changed ({kinds.count('erase')} erased, " \
              f"{kinds.count('text')} retyped, {kinds.count('image')} replaced with an image)."
    if notes:
        summary += " " + " · ".join(dict.fromkeys(notes))
    return jsonify(ok=True,
                   file=register(p, f"{stem(info['name'])}_edited.pdf", "pdf", "output"),
                   note=summary)


@app.route("/api/replace", methods=["POST"])
def replace_text():
    """Find a string and set the replacement in the same font, size and colour."""
    d = request.get_json(force=True)
    info = get_file(d.get("id"), "pdf")
    find = d.get("find") or ""
    new = d.get("replace")
    if not find.strip():
        return err("Enter the text to find.")
    if new is None:
        return err("Enter the replacement text.")
    doc = open_doc(info)
    pages = parse_pages(d.get("pages", "all"), doc.page_count)
    match_case = bool(d.get("match_case"))
    # Arabic PDFs store shaped glyphs, and different producers store them in
    # different orders, so what the user typed matches nothing on its own.
    # Try the plain text, the shaped text, and the shaped-and-reordered text.
    queries = [find]
    if has_arabic(find):
        try:
            import arabic_reshaper
            queries.append(arabic_reshaper.reshape(find))
        except Exception:
            pass
        queries.append(shape_ar(find))
    queries = list(dict.fromkeys(q for q in queries if q))
    hits, notes = 0, []
    try:
        for pn in pages:
            page = doc[pn]
            # search_for ignores case; when the caller wants case-sensitivity,
            # keep only the hits whose actual text matches exactly.
            found = []
            for q in queries:
                for r in editing.merge_glyph_hits(page, page.search_for(q) or [], q):
                    if not any(abs(r & e) > 0 for e in found):
                        found.append(r)
            if match_case:
                wanted = {q.strip() for q in queries}
                found = [r for r in found if page.get_textbox(r).strip() in wanted]
            if not found:
                continue
            spans = [editing.span_at(page, r) for r in found]
            for r in found:
                page.add_redact_annot(r)
            page.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_NONE,
                                  graphics=pymupdf.PDF_REDACT_LINE_ART_NONE,
                                  text=pymupdf.PDF_REDACT_TEXT_REMOVE)
            for r, span in zip(found, spans):
                span = span or {}
                size = float(span.get("size") or 11)
                rgb = span.get("rgb", (0, 0, 0))
                arabic = has_arabic(new)
                choice = editing.resolve_font(doc, page, span, shape_ar(new),
                                              arabic_font(), arabic)
                cleanup = choice.pop("cleanup", None)
                # Keep the replacement inside the space the original occupied —
                # a PDF does not reflow, so anything wider would simply print
                # over the words next to it. Anchor it at the edge the original
                # started from: the right for Arabic, the left otherwise.
                slot = max(r.width, size * 0.6)
                fitted, overflow = editing.size_to_fit(choice, shape_ar(new), size, slot)
                room = max(r.width, fitted * len(new) * 0.85)
                if arabic:
                    box = pymupdf.Rect(min(r.x0, r.x1 - room), r.y0, r.x1, r.y1)
                    align = 2
                else:
                    box = pymupdf.Rect(r.x0, r.y0, max(r.x1, r.x0 + room), r.y1)
                    align = 0
                used = editing.write_into(page, box, shape_ar(new), choice, fitted, rgb,
                                          align=align)
                if overflow:
                    notes.append("the replacement is much longer than the original and "
                                 "may reach the text beside it")
                elif used and abs(used - size) > 0.3:
                    notes.append(f"reduced from {size:g}pt to {used:g}pt to fit the space")
                if cleanup:
                    try:
                        os.remove(cleanup)
                    except OSError:
                        pass
                if used is not None:
                    hits += 1
                    notes.append(choice["used"])
    except Exception as e:
        doc.close()
        return err(f"Replace failed: {e}")

    if not hits:
        doc.close()
        return err(f"“{find}” was not found on the selected pages.")
    p = out_path(f"{stem(info['name'])}_replaced.pdf")
    doc.save(p, garbage=3, deflate=True)
    doc.close()
    note = f"{hits} occurrence(s) replaced. Set in: " + " · ".join(dict.fromkeys(notes))
    return jsonify(ok=True,
                   file=register(p, f"{stem(info['name'])}_replaced.pdf", "pdf", "output"),
                   note=note)


@app.route("/api/page_images/<fid>/<int:page_no>")
def page_images(fid, page_no):
    """Every image on the page, with where it sits — for removing a logo."""
    info = get_file(fid, "pdf")
    doc = open_doc(info)
    if not (0 <= page_no < doc.page_count):
        doc.close()
        abort(400, description="Page number out of range.")
    page = doc[page_no]
    pw, ph = page.rect.width, page.rect.height
    items = []
    for entry in page.get_images(full=True):
        xref = entry[0]
        for r in page.get_image_rects(xref) or []:
            items.append({
                "xref": xref,
                "rect": {"x0": r.x0 / pw, "y0": r.y0 / ph,
                         "x1": r.x1 / pw, "y1": r.y1 / ph},
                "w": round(r.width), "h": round(r.height),
            })
    doc.close()
    return jsonify(ok=True, images=items, count=len(items))


# ============================================================================
# 11) OCR
# ============================================================================
@app.route("/api/ocr", methods=["POST"])
def ocr():
    d = request.get_json(force=True)
    info = get_file(d.get("id"), "pdf")
    lang = d.get("lang", "ara+eng")
    try:
        doc = pymupdf.open(info["path"])
        if doc.needs_pass:
            doc.authenticate("")
        out = pymupdf.open()
        for pg in doc:
            pix = pg.get_pixmap(dpi=int(d.get("dpi", 200)))
            out.insert_pdf(pymupdf.open("pdf", pix.pdfocr_tobytes(language=lang)))
        p = out_path(f"{stem(info['name'])}_ocr.pdf")
        out.save(p, garbage=3, deflate=True)
        out.close()
        doc.close()
        return jsonify(ok=True,
                       file=register(p, f"{stem(info['name'])}_ocr.pdf", "pdf", "output"),
                       note=f"Searchable copy created (language: {lang}).")
    except Exception as e:
        return err("OCR needs Tesseract installed on this machine — including the "
                   f"Arabic language pack for 'ara'. Details: {e}")


# ============================================================================
# 12) Updates
# ============================================================================
import updater


@app.route("/api/update/check", methods=["POST", "GET"])
def update_check():
    return jsonify(**updater.check())


@app.route("/api/update/apply", methods=["POST"])
def update_apply():
    return jsonify(**updater.apply())


@app.route("/api/update/rollback", methods=["POST"])
def update_rollback():
    return jsonify(**updater.rollback())


@app.route("/api/health")
def health():
    fp = arabic_font()
    return jsonify(ok=True, app_version=updater.current_version(),
                   version=pypdf.__version__, files=len(FILES),
                   arabic_font=os.path.basename(fp) if fp else None,
                   bundled=fp == BUNDLED_REGULAR,
                   arabic_shaping=SHAPING_OK,
                   arabic_shaping_error=SHAPING_ERROR)


# ============================================================================
def main():
    host = os.environ.get("PDFSTUDIO_HOST", "127.0.0.1")
    port = int(os.environ.get("PDFSTUDIO_PORT", "8765"))
    url = f"http://{host}:{port}"
    fp = arabic_font()
    print("\n" + "=" * 60)
    print("  PDF Studio — local PDF toolkit")
    print(f"  Open your browser at:  {url}")
    print(f"  Arabic font: {os.path.basename(fp) if fp else 'NOT FOUND — see fonts/README.txt'}")
    print(f"  Arabic shaping: {'OK' if SHAPING_OK else 'UNAVAILABLE — ' + SHAPING_ERROR}")
    print("  Press Ctrl+C to stop.")
    print("=" * 60 + "\n")
    if os.environ.get("PDFSTUDIO_NOBROWSER") != "1":
        threading.Timer(1.2, lambda: webbrowser.open(url)).start()
    app.run(host=host, port=port, debug=False, threaded=True)


if __name__ == "__main__":
    main()
